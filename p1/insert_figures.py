"""
Place the figure PNGs into the Word file, each after the paragraph that cites it.

Reads paper_v17_zotero.docx (the version with live citation fields) and writes
paper_v17_submission.docx, the file to submit.

Each figure is inserted centred, with a caption below it in the usual style.
If a figure needs changing later, edit figures_v14.xlsx or the PowerPoint,
re-run export_figures_png.py, and swap that one image in Word by hand.

    python p1/insert_figures.py
"""
import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt, RGBColor

HERE = Path(__file__).resolve().parent
EXP = HERE.parent
ROOT = EXP.parent
PNG = EXP / "figures_png"
SRC = ROOT / "paper_v17_zotero.docx"
OUT = ROOT / "paper_v17_submission.docx"

# figure -> (anchor: a distinctive phrase in the paragraph it belongs after, caption)
FIGURES = {
    1: ("Figure 1 shows how the parts fit together",
        "Figure 1. The experimental pipeline and the four-part evaluation protocol. "
        "Parts 1 to 3 test whether the predicted number carries prognostic information; "
        "part 4 tests whether the explanation describes the input it was given."),
    2: ("Figure 2 plots predicted against true RUL",
        "Figure 2. Predicted against true RUL on FD001. Points on the dashed diagonal would be "
        "exact. The prompted model forms horizontal bands: its output varies little with the "
        "engine, and the bands themselves show that the dependence is weak rather than "
        "absent."),
    3: ("The anchor is model-specific",
        "Figure 3. Distribution of predicted RUL on FD001 for the four models run on all 100 "
        "test engines at the canonical configuration, with XGBoost and the true RUL for "
        "reference. DeepSeek-R1, run on 50 engines, is reported in Table 3. A regressor "
        "spreads across the range; the prompted models do not."),
    4: ("The same integer comes out of all of them",
        "Figure 4. Share of predictions falling on the single most frequent value, on "
        "each of the four C-MAPSS sub-datasets, zero-shot with n = 30. The label on each "
        "bar gives the number of distinct values the model produced over the engines "
        "named in the category label. For comparison, the supervised regressors on FD001 and "
        "FD003 occupy 59 to 69 values on a one-cycle grid, with no value taking more than "
        "6% of the predictions."),
    5: ("Figure 5 breaks these results down by sensor",
        "Figure 5. What the stated direction agrees with, per sensor: the reference direction, the empirical benchmark direction, the direction in the supplied window, and the base rate on the same claims. Cued sensors are the ones the prompt names. The benchmark series is computed only where the sub-dataset attests a direction, which excludes s7, s12 and s15 on FD003 and s14 on FD001 and leaves 2,754 of the 3,816 claims; the other three series use all of them. Where the two references are opposite, at s9, s12 and s14, the claims follow the reference. Faithfulness and the base rate track each other everywhere. Bars rest on between 10 (s4) and 947 (s11) directional claims."),
    6: ("Reversing the series left the stated directions where they were",
        "Figure 6. The trend-reversal test, with each bar labelled and the share shown on the changed-claim series, since 6 pairs against 253 is otherwise invisible. For each sensor, the number of (engine, sensor) "
        "pairs whose trend genuinely reversed between the two arms, and the number of those "
        "in which the direction stated by the model reversed with it."),
    7: ("Figure 7 plots that relation over the seven example sets",
        "Figure 7. Mean predicted RUL against the mean RUL of the examples placed in the "
        "prompt, over seven example sets. The dashed line is the test-set mean."),
    8: ("Figure 8 reports the rank correlation of every FD001 configuration",
        "Figure 8. Spearman rank correlation between predicted and true RUL for the 15 FD001 "
        "runs of the main grid, with 95% bootstrap intervals. Random Forest reaches +0.813 "
        "and the repaired LSTM +0.889 on the same engines."),
}


# Alternative text, one entry per figure number. Read by screen readers and
# checked by several publishers' accessibility tools; it should describe what the
# figure shows, not repeat the caption.
ALT_TEXT = {
    1: "Flow diagram of the evaluation protocol: the C-MAPSS sub-datasets feed both the "
       "prompted language models and the supervised baselines, whose outputs are scored on "
       "error, rank correlation, output entropy and explanation faithfulness.",
    2: "Predicted against true remaining useful life for the reference configuration. The "
       "predictions lie on a small number of horizontal lines instead of following the "
       "diagonal.",
    3: "Distribution of predicted remaining useful life on FD001, one series per model, "
       "showing the prompted models concentrated in one or two bins while the reference "
       "distribution spreads across the range.",
    4: "Bar chart of the share of predictions falling on a single value for each of the "
       "four C-MAPSS sub-datasets, between 51% and 83%, each bar labelled with the number "
       "of distinct values produced.",
    5: "Per sensor, the share of directional claims agreeing with the reference direction, "
       "with the direction measured in the benchmark, with the window shown, and the "
       "base rate on the same claims.",
    6: "Counts of paired responses in which the stated sensor direction did or did not follow "
       "the reversal applied to the input trend.",
    7: "Mean predicted remaining useful life against the mean remaining useful life of the "
       "few-shot examples, one point per example set, with a fitted line.",
    8: "Rank correlation with bootstrap intervals for the 15 FD001 runs of the main grid, "
       "against the supervised baselines on the same engines.",
}

def add_figure(doc_body, after_par, img, caption, width_cm=15.5, alt=None, num=None):
    """Insert a picture paragraph and a caption paragraph after a given one."""
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement

    pic_p = OxmlElement("w:p")
    after_par._p.addnext(pic_p)
    from docx.text.paragraph import Paragraph
    pic_par = Paragraph(pic_p, after_par._parent)
    pic_par.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pic_par.add_run().add_picture(str(img), width=Cm(width_cm))
    if alt:
        for dp in pic_p.findall(".//" + qn("wp:docPr")):
            dp.set("descr", alt)
            if num:
                dp.set("title", f"Figure {num}")
    # a figure should not be separated from its caption
    pPr = pic_p.get_or_add_pPr()
    if pPr.find(qn("w:keepNext")) is None:
        pPr.append(OxmlElement("w:keepNext"))

    cap_p = OxmlElement("w:p")
    pic_p.addnext(cap_p)
    cap_par = Paragraph(cap_p, after_par._parent)
    cap_par.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    r = cap_par.add_run(caption)
    r.font.size = Pt(9)
    r.font.color.rgb = RGBColor(0x40, 0x40, 0x40)
    return cap_par


def main():
    if not SRC.exists():
        print(f"missing {SRC.name}: run insert_zotero_fields.py first")
        return 1
    doc = Document(SRC)
    placed, missing = [], []

    for num, (anchor, caption) in sorted(FIGURES.items()):
        img = PNG / f"fig{num}.png"
        if not img.exists():
            missing.append(f"fig{num}.png not found")
            continue
        target = None
        for par in doc.paragraphs:
            if anchor.lower() in par.text.lower():
                target = par
                break
        if target is None:
            missing.append(f"figure {num}: anchor not found ({anchor[:40]}...)")
            continue
        add_figure(doc.element.body, target, img, caption,
                   alt=ALT_TEXT.get(num), num=num)
        placed.append(num)

    # Word shows these in File > Info; python-docx leaves its own name there,
    # which should not go to an editor.
    cp = doc.core_properties
    cp.author = "Filippo De Carlo"
    cp.last_modified_by = "Filippo De Carlo"
    cp.title = ("Evaluating large language models for remaining useful life prediction: "
                "a protocol and evidence from C‑MAPSS")
    cp.comments = ""
    cp.category = ""
    cp.keywords = ("large language models; remaining useful life; prognostics and health "
                   "management; evaluation protocol; explanation faithfulness")

    doc.save(OUT)
    print(f"[saved] {OUT}")
    print(f"figures placed: {placed}")
    for m in missing:
        print(f"  ! {m}")
    ga = PNG / "graphical_abstract.png"
    if ga.exists():
        print(f"\nGraphical abstract is a separate upload: {ga}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
