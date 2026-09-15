"""
P3.3c - Turn the author's reading into the labels file the validation scores.

The reading was done in a small web app rather than in the Excel sheet built by
p3_3_author_sheet.py: the same 139 cases in the same shuffled order, the full
response with every occurrence of the sensor's name highlighted, a chart of the
values the model had in its prompt, and the six labels as buttons. The app keeps,
for each case, the first label given (`first_label`) and whether the provisional
label that preceded this reading was opened afterwards (`revealed`). The first
label is the independent reading and is the one used.

Either source is accepted, so a replication can use the sheet instead:

    python p1/p3_3_import_readings.py                       # the app export
    python p1/p3_3_import_readings.py path/to/export.json
    python p1/p3_3_import_readings.py path/to/sheet.xlsx

Output: results_p1/p3_3_author_labels.csv, one row per sampled pair, keyed by id.
"""
import json
import sys
from pathlib import Path

import pandas as pd

HERE = Path(__file__).resolve().parent
EXP = HERE.parent
ROOT = EXP.parent
P1 = EXP / "results_p1"
KEY = P1 / "p3_3_author_key.csv"
OUT = P1 / "p3_3_author_labels.csv"
DEFAULT = P1 / "p3_3_author_reading_export.json"
VALID = {"increase", "decrease", "stable", "unspecified", "absent", "ambiguous"}


def from_app(path):
    recs = json.loads(Path(path).read_text(encoding="utf-8"))
    if isinstance(recs, dict):
        recs = recs.get("readings", recs.get("letture"))
    recs = [r.get("data", r) for r in recs]
    df = pd.DataFrame(recs)
    return pd.DataFrame({"case": df["case"].astype(int), "sensor_read": df["sensor"],
                         "author": df["first_label"].fillna("").str.strip().str.lower(),
                         "author_note": df.get("note", pd.Series([""] * len(df))).fillna(""),
                         "revealed": df.get("revealed", pd.Series([False] * len(df)))
                         .fillna(False).astype(bool)})


def from_sheet(path):
    import openpyxl
    ws = openpyxl.load_workbook(path, data_only=True)["Casi"]
    head = [str(c.value or "").strip() for c in ws[1]]
    need = {"case": "Caso", "sensor": "Sensore", "label": "La tua lettura", "note": "Note"}
    missing = [h for h in need.values() if h not in head]
    if missing:
        raise SystemExit(f"the sheet has no column {missing}")
    col = {k: head.index(h) for k, h in need.items()}
    rows = [{"case": int(r[col["case"]]), "sensor_read": r[col["sensor"]],
             "author": str(r[col["label"]] or "").strip().lower(),
             "author_note": str(r[col["note"]] or "").strip(), "revealed": False}
            for r in ws.iter_rows(min_row=2, values_only=True) if r[col["case"]] is not None]
    return pd.DataFrame(rows)


def main():
    src = Path(sys.argv[1]) if len(sys.argv) > 1 else DEFAULT
    lab = from_sheet(src) if src.suffix.lower() == ".xlsx" else from_app(src)
    key = pd.read_csv(KEY)
    if lab.case.duplicated().any():
        raise SystemExit(f"cases read twice: {lab[lab.case.duplicated()].case.tolist()}")
    lab = key.merge(lab, on="case", how="left")
    empty = lab[lab.author.fillna("") == ""].case.tolist()
    if empty:
        raise SystemExit(f"The reading is not complete: {len(empty)} of {len(key)} cases have "
                         f"no label (cases {', '.join(map(str, empty[:15]))}"
                         f"{' ...' if len(empty) > 15 else ''}).")
    bad = lab[~lab.author.isin(VALID)]
    if len(bad):
        raise SystemExit(f"labels outside the six allowed ones: "
                         f"{bad[['case', 'author']].to_dict('records')}")
    sample = pd.read_csv(P1 / "p3_3_extractor_sample.csv", usecols=["id", "sensor"])
    chk = lab.merge(sample, on="id")
    if (chk.sensor != chk.sensor_read).any():
        raise SystemExit("the reading and the sample disagree on the sensor for ids "
                         f"{chk[chk.sensor != chk.sensor_read].id.tolist()}")
    out = chk[["id", "case", "sensor", "author", "author_note", "revealed"]].sort_values("id")
    out.to_csv(OUT, index=False, encoding="utf-8")
    print(f"{len(out)} labels from {src.name}; provisional label opened after the reading "
          f"in {int(out.revealed.sum())} cases")
    print(out.author.value_counts().to_string())
    print(f"[saved] {OUT}")


if __name__ == "__main__":
    main()
